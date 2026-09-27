"""
提取当日 Excel 日报为纯文本，供公众号文章生成 / 自动化使用。
用法:
    python utils/extract_report.py            # 默认今天
    python utils/extract_report.py 2026-07-23
输出到 stdout，并同时保存到 logs/wechat/report_summary_YYYYMMDD.txt
"""
import sys
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent.parent
REPORT_DIR = BASE / "logs" / "reports"


def extract(date_str: str = None) -> str | None:
    if date_str is None:
        date_str = datetime.now().strftime("%Y-%m-%d")
    # 日报文件名用 YYYYMMDD
    ymd = date_str.replace("-", "")
    xlsx = REPORT_DIR / f"日报_{ymd}.xlsx"
    if not xlsx.exists():
        print(f"[WARN] 未找到日报文件: {xlsx}（可能当日无交易/休市）")
        return None

    wb = load_workbook(xlsx, data_only=True)
    lines = [f"# ETF交易日报 - {date_str}", ""]
    for ws in wb.worksheets:
        lines.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if all(c == "" for c in cells):
                continue
            line = " | ".join(cells).strip(" |")
            lines.append(line)
        lines.append("")

    text = "\n".join(lines)

    out_dir = BASE / "logs" / "wechat"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_file = out_dir / f"report_summary_{ymd}.txt"
    out_file.write_text(text, encoding="utf-8")

    print(text)
    print(f"\n[OK] 文本摘要已保存: {out_file}")
    return text


if __name__ == "__main__":
    d = sys.argv[1] if len(sys.argv) > 1 else None
    extract(d)
